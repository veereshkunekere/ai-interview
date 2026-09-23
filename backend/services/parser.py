"""Extract plain text from uploaded resume / job-description files."""

from docx import Document
from pypdf import PdfReader


def extract_text(filename: str, data: bytes) -> str:
    """Return the plain text content of a PDF, DOCX, or TXT file.

    Raises ValueError when the file cannot be parsed (e.g. a scanned PDF
    with no extractable text layer).
    """
    name = (filename or "").lower()

    if name.endswith(".pdf") or data[:4] == b"%PDF":
        return _from_pdf(data)

    if name.endswith(".docx"):
        return _from_docx(data)

    # Plain text (and unknown-but-text formats) as a fallback.
    text = data.decode("utf-8", errors="ignore")
    if not text.strip():
        raise ValueError("File is empty or contains no readable text.")
    return text


def _from_pdf(data: bytes) -> str:
    import io

    reader = PdfReader(io.BytesIO(data))
    parts = [(page.extract_text() or "") for page in reader.pages]
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError(
            "No extractable text found in this PDF. It may be a scanned "
            "image PDF - please upload a text-based PDF, DOCX, or TXT file."
        )
    return text


def _from_docx(data: bytes) -> str:
    import io

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("DOCX file contains no readable text.")
    return text
